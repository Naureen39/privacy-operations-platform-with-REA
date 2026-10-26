"""Phase 11 acceptance: a full system onboarding and a full access request, run from API calls only.

Nothing here touches the database or an internal function: every step is an HTTP request made by
the person who would make it (analyst, privacy officer), including reading the progress stream."""

import io
import json
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.records.builder import NOTICE
from tests.fixtures_api import client_for, csrf
from tests.fixtures_fusion import DATASETS, ESTATE, FusionEstate

ALL_CHECKS = {
    "name_matches_records": True,
    "contact_details_confirmed": True,
    "identity_document_seen": True,
}


def upload(client: TestClient, system_id: str, kind: str, path: Any) -> None:
    response = client.post(
        f"/api/v1/systems/{system_id}/evidence",
        data={"kind": kind},
        files={"file": (path.name, path.read_bytes(), "application/octet-stream")},
        headers=csrf(client),
    )
    assert response.status_code == 201, response.text


def stream(client: TestClient, url: str) -> list[tuple[str, dict[str, Any]]]:
    events: list[tuple[str, dict[str, Any]]] = []
    name = ""
    with client.stream("GET", url) as response:
        for line in response.iter_lines():
            if line.startswith("event: "):
                name = line.removeprefix("event: ")
            elif line.startswith("data: "):
                events.append((name, json.loads(line.removeprefix("data: "))))
    return events


def test_a_system_is_onboarded_from_evidence_to_an_approved_exported_record(
    fusion_estate: FusionEstate,
) -> None:
    fe = fusion_estate
    analyst = client_for(fe.app, fe.users["analyst"])
    officer = client_for(fe.app, fe.users["admin"])
    estate = ESTATE / "clinic_portal"

    # 1. the analyst registers the system and attaches the evidence
    system = analyst.post(
        "/api/v1/systems",
        json={
            "name": "Clinic end to end",
            "description": "Appointment booking for the clinic",
            "owner_user_id": str(fe.users["system_owner"].id),
        },
        headers=csrf(analyst),
    )
    assert system.status_code == 201
    system_id = system.json()["id"]
    for kind, name in (
        ("har", "evidence/capture.har"),
        ("website", "evidence/website_snapshot.html"),
        ("javascript_bundle", "evidence/app.bundle.js"),
        ("sqlite_snapshot", "evidence/snapshot.sqlite"),
        ("privacy_policy", "policy.md"),
        ("questionnaire", "questionnaire.json"),
    ):
        upload(analyst, system_id, kind, estate / name)
    assert len(analyst.get(f"/api/v1/systems/{system_id}/evidence").json()) == 6

    # 2. onboarding runs as a job; the progress arrives as events
    job = analyst.post(
        f"/api/v1/systems/{system_id}/onboard",
        json={"scale": {"data_subjects_estimate": 284}},
        headers=csrf(analyst),
    )
    assert job.status_code == 202
    job_id = job.json()["id"]
    events = stream(analyst, f"/api/v1/jobs/{job_id}/events")
    names = [n for n, _ in events]
    assert (
        names[0] == "queued"
        and names[-2:] == ["succeeded", "end"]
        and names.count("stage_succeeded") == 5
    )
    assert analyst.get(f"/api/v1/jobs/{job_id}").json()["state"] == "succeeded"

    # 3. what the job produced
    elements = analyst.get(
        f"/api/v1/systems/{system_id}/data-elements", params={"limit": 1000}
    ).json()
    assert len(elements) > 20 and any(e["special_category"] for e in elements)
    assert {r["domain"] for r in analyst.get(f"/api/v1/systems/{system_id}/recipients").json()} >= {
        "cal.bookwell.example"
    }
    findings = analyst.get(
        "/api/v1/findings", params={"system_id": system_id, "status": "open"}
    ).json()
    types = {f["type"] for f in findings}
    assert {"missing_art9_condition", "third_country_no_safeguard", "stale_review"} <= types
    record = next(
        a for a in analyst.get("/api/v1/activities").json() if system_id in a["system_ids"]
    )
    assert record["status"] == "draft" and record["missing_fields"] == [
        "special_category_condition"
    ]
    screening = analyst.get(f"/api/v1/activities/{record['id']}/dpia-screening").json()
    assert screening["required"] is True and {
        c["id"] for c in screening["criteria"] if c["met"]
    } == {"sensitive_data", "vulnerable_subjects"}

    # 4. the analyst turns a finding into a task, records the missing condition and submits
    task = analyst.post(
        f"/api/v1/findings/{next(f['id'] for f in findings if f['type'] == 'missing_art9_condition')}/task",
        headers=csrf(analyst),
    )
    assert task.status_code == 201
    fixed = analyst.patch(
        f"/api/v1/activities/{record['id']}",
        json={"special_category_condition": "Art. 9(2)(h) health care"},
        headers=csrf(analyst),
    )
    assert fixed.json()["missing_fields"] == []
    analyst.patch(
        f"/api/v1/tasks/{task.json()['id']}", json={"status": "done"}, headers=csrf(analyst)
    )
    submitted = analyst.post(f"/api/v1/activities/{record['id']}/submit", headers=csrf(analyst))
    assert submitted.status_code == 200 and submitted.json()["status"] == "awaiting_review"

    # 5. the privacy officer is told, approves and publishes
    assert any(
        n["kind"] == "approval_requested" and n["entity_id"] == record["id"]
        for n in officer.get("/api/v1/notifications").json()
    )
    assert (
        officer.post(
            f"/api/v1/activities/{record['id']}/approve",
            json={"note": "Checked against the evidence."},
            headers=csrf(officer),
        ).json()["status"]
        == "approved"
    )
    assert (
        officer.post(f"/api/v1/activities/{record['id']}/publish", headers=csrf(officer)).json()[
            "status"
        ]
        == "published"
    )
    assert any(n["kind"] == "approval_decided" for n in analyst.get("/api/v1/notifications").json())

    # 6. the impact assessment is drafted, completed and exported
    dpia = officer.post(f"/api/v1/activities/{record['id']}/dpia", json={}, headers=csrf(officer))
    assert dpia.status_code == 201 and dpia.json()["missing_fields"] == []
    pdf = officer.get(f"/api/v1/dpias/{dpia.json()['id']}/export.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")

    # 7. the register export carries the record
    workbook = load_workbook(io.BytesIO(officer.get("/api/v1/exports/register.xlsx").content))
    assert workbook["Register"]["A2"].value == NOTICE
    assert "Clinic end to end" in {
        row[0].value for row in workbook["Register"].iter_rows(min_row=6) if row[0].value
    }

    # 8. a second run changes nothing
    again = analyst.post(
        f"/api/v1/systems/{system_id}/onboard",
        json={"scale": {"data_subjects_estimate": 284}},
        headers=csrf(analyst),
    ).json()
    assert again["state"] == "succeeded" and again["result"]["record"]["changed"] is False


def test_an_access_request_is_handled_from_intake_to_a_closed_case(
    fusion_estate: FusionEstate,
) -> None:
    fe = fusion_estate
    analyst = client_for(fe.app, fe.users["analyst"])
    officer = client_for(fe.app, fe.users["admin"])
    case = next(
        c
        for c in json.loads(
            (DATASETS / "people" / "request_cases.json").read_text(encoding="utf-8")
        )["cases"]
        if c["id"] == "RC04"
    )
    email = case["identifier"]["value"]
    base = "/api/v1/access-requests"

    request = analyst.post(
        base, json={"identifier": email, "received_date": "2026-09-20"}, headers=csrf(analyst)
    ).json()
    assert (
        request["status"] == "received"
        and request["deadline_date"] == "2026-10-20"
        and request["reminder"] == "none"
    )
    rid = request["id"]
    assert (
        analyst.post(
            f"{base}/{rid}/identity", json={"checklist": ALL_CHECKS}, headers=csrf(analyst)
        ).json()["identity_verified"]
        is True
    )
    found = analyst.post(
        f"{base}/{rid}/search", json={"identifier": email}, headers=csrf(analyst)
    ).json()
    by_name = {}
    for r in found["results"]:
        by_name.setdefault(r["system_name"], set()).add(r["table"])
    for system in case["expected_systems"]:
        assert system in by_name, system
    expected_locations = {(loc["system"], loc["table"]) for loc in case["expected_locations"]}
    assert expected_locations <= {(s, t) for s, tables in by_name.items() for t in tables}

    letter = analyst.post(
        f"{base}/{rid}/response", json={"cover_note": True}, headers=csrf(analyst)
    ).json()
    assert (
        letter["body"].startswith(NOTICE) and email not in letter["body"] and letter["cover_note"]
    )
    assert (
        analyst.post(f"{base}/{rid}/submit", headers=csrf(analyst)).json()["status"]
        == "awaiting_approval"
    )
    assert any(
        n["kind"] == "approval_requested" and n["entity_id"] == rid
        for n in officer.get("/api/v1/notifications").json()
    )
    assert officer.post(f"{base}/{rid}/approve", json={}, headers=csrf(officer)).status_code == 200
    assert (
        officer.post(f"{base}/{rid}/respond", headers=csrf(officer)).json()["status"] == "responded"
    )
    closed = analyst.post(f"{base}/{rid}/close", headers=csrf(analyst)).json()
    assert closed["status"] == "closed" and closed["closed_at"] and closed["reminder"] == "none"
    assert rid not in {r["id"] for r in analyst.get(f"{base}/reminders").json()}
