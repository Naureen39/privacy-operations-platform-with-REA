"""Records of processing: generation from the reference estate, citations, versions and diffs, the
review workflow and the exports."""

import io
import json
from typing import Any

import pdfplumber
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, text

from app.records.builder import NOTICE, REQUIRED_FIELDS, completeness
from tests.fixtures_api import client_for, csrf
from tests.fixtures_fusion import ESTATE, FusionEstate
from tests.test_fusion_acceptance import fuse, pipeline


@pytest.fixture(scope="module")
def records(estate_generation: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """The generated record of every estate system (the answer of the first generation)."""
    return estate_generation


def record_for(
    fusion_estate: FusionEstate, client: TestClient, estate_id: str, name: str, owner: Any = None
) -> tuple[str, dict[str, Any]]:
    system_id = pipeline(client, estate_id, owner or fusion_estate.users["system_owner"].id, name)
    fuse(client, system_id)
    response = client.post(f"/api/v1/systems/{system_id}/record", headers=csrf(client))
    assert response.status_code == 200, response.text
    return system_id, response.json()["activity"]


# -- generation ----------------------------------------------------------------------------------


def test_every_estate_system_gets_a_record_complete_against_the_required_fields(
    records: dict[str, dict[str, Any]],
) -> None:
    assert len(records) == 8
    gaps = {}
    for estate_id, result in records.items():
        activity = result["activity"]
        assert (
            result["created"] is True
            and activity["status"] == "draft"
            and activity["current_version"] == 1
        )
        for name in REQUIRED_FIELDS:
            if name in ("recipients", "transfers"):
                assert activity[name] is not None
            else:
                assert activity[name] not in (None, "", []), (estate_id, name)
        if activity["missing_fields"]:
            gaps[estate_id] = activity["missing_fields"]
    # the only gaps are real ones: special category data with no Article 9 condition recorded
    assert gaps == {
        "clinic_portal": ["special_category_condition"],
        "hr_onboarding": ["special_category_condition"],
    }


def test_every_field_cites_evidence_that_exists(
    records: dict[str, dict[str, Any]], fusion_estate: FusionEstate, owner_engine: Engine
) -> None:
    with owner_engine.connect() as conn:
        org = {"o": fusion_estate.org_id}
        discovery = {
            str(r[0])
            for r in conn.execute(text("SELECT id FROM discovery_evidence WHERE org_id = :o"), org)
        }
        evidence = {
            str(r[0])
            for r in conn.execute(text("SELECT id FROM evidence_items WHERE org_id = :o"), org)
        }
        systems = {
            str(r[0]) for r in conn.execute(text("SELECT id FROM systems WHERE org_id = :o"), org)
        }
    known = {
        "discovery": discovery,
        "questionnaire": evidence,
        "policy": evidence,
        "system": systems,
    }
    checked = 0
    for estate_id, result in records.items():
        activity = result["activity"]
        for name in (*REQUIRED_FIELDS, "special_category_condition"):
            refs = activity["citations"].get(name)
            assert refs, (estate_id, name)
            for ref in refs:
                assert ref["kind"] in {*known, "assistant", "manual"} and ref["detail"], (
                    estate_id,
                    name,
                    ref,
                )
                if ref["kind"] in known:
                    assert ref["id"] in known[ref["kind"]], (estate_id, name, ref)
                checked += 1
    assert checked > 200


def test_record_content_comes_from_the_facts(records: dict[str, dict[str, Any]]) -> None:
    shop = records["web_shop"]["activity"]
    assert shop["data_subjects"] == ["customers"] and shop["lawful_basis"] == "contract"
    vendors = {r["vendor"]: r for r in shop["recipients"]}
    assert {"Stripe", "Meta Platforms"} <= set(vendors)
    assert (
        vendors["Stripe"]["role"] == "processor"
        and vendors["Stripe"]["country"] == "US"
        and vendors["Stripe"]["declared"] is True
    )
    assert (
        vendors["Meta Platforms"]["declared"] is False
        and "not declared" in vendors["Meta Platforms"]["text"]
    )
    transfers = {t["recipient"]: t for t in shop["transfers"]}
    assert (
        transfers["Stripe"]["mechanism"] == "standard contractual clauses"
        and transfers["Stripe"]["gap"] is False
    )
    assert (
        transfers["Meta Platforms"]["gap"] is True
        and "no safeguard recorded" in transfers["Meta Platforms"]["text"]
    )
    assert "7 years" in shop["retention"] and "date_of_birth" in shop["data_categories"]
    assert "Encryption" in shop["security_measures"] or shop["security_measures"].startswith("- ")
    assert shop["next_review_date"] == "2027-03-12"  # twelve months after the last review

    clinic = records["clinic_portal"]["activity"]
    assert clinic["special_categories"] and "health_data" in clinic["special_categories"]
    assert "special category data" in clinic["summary"] and "Health data" in clinic["summary"]
    assert (
        clinic["next_review_date"] == "2026-11-01"
    )  # overdue: a review within a month of the comparison date
    uk = {t["recipient"]: t["mechanism"] for t in records["web_analytics"]["activity"]["transfers"]}
    assert uk["BrightLens Agency"] == "adequacy decision"
    assert records["newsletter"]["activity"]["transfers"][0]["country"] == "US"


def test_internal_hosts_and_eea_recipients_are_not_transfers(
    records: dict[str, dict[str, Any]],
) -> None:
    helpdesk = records["helpdesk"]["activity"]
    assert helpdesk["transfers"] == []
    assert all(r["country"] != "US" for r in helpdesk["recipients"])


def test_purposes_and_basis_are_suggestions_marked_for_review(
    records: dict[str, dict[str, Any]], fusion_estate: FusionEstate
) -> None:
    shop = records["web_shop"]
    assert shop["suggestion_used"] is True and 0 < shop["tokens_used"] < 1500
    refs = shop["activity"]["citations"]["purposes"]
    assert any(
        r["kind"] == "assistant" and r["detail"].startswith("Suggested for review") for r in refs
    )
    assert any(r["kind"] == "questionnaire" for r in refs)  # the owner's own words stay
    declared = json.loads((ESTATE / "web_shop" / "questionnaire.json").read_text(encoding="utf-8"))[
        "answers"
    ]["purposes"]
    assert set(declared) <= set(shop["activity"]["purposes"])
    total = sum(r["tokens_used"] for r in records.values())
    assert (
        total / 8 < 1000
    )  # one purpose and basis call per system, well inside the 4,000 token budget


def test_the_assistant_never_overrides_the_owners_stated_basis(
    records: dict[str, dict[str, Any]],
) -> None:
    assert (
        records["clinic_portal"]["activity"]["lawful_basis"] == "legitimate_interests"
    )  # as the owner wrote it
    refs = records["clinic_portal"]["activity"]["citations"]["lawful_basis"]
    assert [r["kind"] for r in refs] == ["questionnaire"]


def test_completeness_check() -> None:
    full = {n: "x" for n in REQUIRED_FIELDS} | {
        "recipients": [],
        "transfers": [],
        "purposes": ["p"],
        "data_subjects": ["s"],
        "data_categories": ["c"],
    }
    assert completeness(full, []) == []
    assert completeness(full, ["health_data"]) == ["special_category_condition"]
    assert completeness({**full, "retention": None, "purposes": []}, []) == [
        "purposes",
        "retention",
    ]
    assert completeness({k: v for k, v in full.items() if k != "recipients"}, []) == ["recipients"]


# -- versions, diffs and edits -------------------------------------------------------------------


def test_regenerating_without_new_evidence_does_not_create_a_version(
    fusion_estate: FusionEstate, records: dict[str, dict[str, Any]]
) -> None:
    admin = fusion_estate.admin
    system_id = fusion_estate.systems["crm"]
    again = admin.post(f"/api/v1/systems/{system_id}/record", headers=csrf(admin)).json()
    assert (
        again["created"] is False
        and again["changed"] is False
        and again["activity"]["current_version"] == 1
    )
    assert (
        again["tokens_used"] == 0 or again["suggestion_used"]
    )  # the model answer came from the cache
    versions = admin.get(f"/api/v1/activities/{again['activity']['id']}/versions").json()
    assert [v["version"] for v in versions] == [1]


def test_new_evidence_creates_a_new_version_with_a_diff(fusion_estate: FusionEstate) -> None:
    admin = fusion_estate.admin
    system_id, activity = record_for(fusion_estate, admin, "newsletter", "Newsletter copy")
    changed = json.loads((ESTATE / "newsletter" / "questionnaire.json").read_text(encoding="utf-8"))
    changed["answers"]["retention"] = "Until unsubscribe plus 14 days"
    changed["answers"]["security_measures"] = [
        *changed["answers"]["security_measures"],
        "Quarterly access review",
    ]
    upload = admin.post(
        f"/api/v1/systems/{system_id}/evidence", data={"kind": "questionnaire"},
        files={"file": ("questionnaire.json", json.dumps(changed).encode(), "application/json")}, headers=csrf(admin),
    )  # fmt: skip
    assert upload.status_code == 201
    fuse(admin, system_id)
    result = admin.post(f"/api/v1/systems/{system_id}/record", headers=csrf(admin)).json()
    assert result["changed"] is True and result["activity"]["current_version"] == 2
    assert "14 days" in result["activity"]["retention"]

    diff = admin.get(
        f"/api/v1/activities/{activity['id']}/diff", params={"from": 1, "to": 2}
    ).json()
    assert diff["from_version"] == 1 and diff["to_version"] == 2
    fields = {c["field"]: c for c in diff["changes"]}
    assert set(fields) == {"retention", "security_measures"}
    assert fields["retention"]["change"] == "changed" and "14 days" in fields["retention"]["after"]
    assert "Quarterly access review" in fields["security_measures"]["after"]
    assert diff["summary"] == "Changed: retention, security measures"
    versions = admin.get(f"/api/v1/activities/{activity['id']}/versions").json()
    assert [v["version"] for v in versions] == [1, 2] and "Regenerated" in versions[1][
        "change_summary"
    ]
    detail = admin.get(f"/api/v1/activities/{activity['id']}/versions/1").json()
    assert detail["snapshot"]["retention"] != result["activity"]["retention"]
    assert admin.get(f"/api/v1/activities/{activity['id']}/versions/9").status_code == 404
    assert (
        admin.get(
            f"/api/v1/activities/{activity['id']}/diff", params={"from": 1, "to": 9}
        ).status_code
        == 404
    )


def test_list_changes_in_a_diff_show_what_was_added_and_removed(
    fusion_estate: FusionEstate,
) -> None:
    admin = fusion_estate.admin
    system_id, activity = record_for(fusion_estate, admin, "staff_desktop", "Desktop for diff")
    edited = admin.patch(
        f"/api/v1/activities/{activity['id']}",
        json={"purposes": ["Payroll", "Time keeping"]},
        headers=csrf(admin),
    )
    assert edited.status_code == 200
    diff = admin.get(
        f"/api/v1/activities/{activity['id']}/diff", params={"from": 1, "to": 2}
    ).json()
    change = diff["changes"][0]
    assert change["field"] == "purposes" and sorted(change["added"]) == ["Payroll", "Time keeping"]
    assert change["removed"] == activity["purposes"] and "+2" in diff["summary"]


def test_a_manual_edit_survives_regeneration_and_is_marked(fusion_estate: FusionEstate) -> None:
    admin = fusion_estate.admin
    system_id, activity = record_for(fusion_estate, admin, "helpdesk", "Helpdesk for edits")
    edited = admin.patch(
        f"/api/v1/activities/{activity['id']}",
        json={"retention": "9 months after closure", "lawful_basis": "contract"},
        headers=csrf(admin),
    )
    assert edited.status_code == 200 and edited.json()["current_version"] == 2
    refs = edited.json()["citations"]["retention"]
    assert refs[-1]["kind"] == "manual" and refs[-1]["id"] == str(fusion_estate.users["admin"].id)
    again = admin.post(f"/api/v1/systems/{system_id}/record", headers=csrf(admin)).json()
    assert again["changed"] is False and again["activity"]["retention"] == "9 months after closure"
    assert again["activity"]["lawful_basis"] == "contract"
    nothing = admin.patch(
        f"/api/v1/activities/{activity['id']}",
        json={"retention": "9 months after closure"},
        headers=csrf(admin),
    )
    assert nothing.json()["current_version"] == 2  # an edit that changes nothing is not a version


@pytest.mark.parametrize(
    "body",
    [
        {"name": "x"},
        {"recipients": []},
        {"purposes": "text"},
        {"next_review_date": "tomorrow"},
        {"unknown": 1},
    ],
)
def test_only_editable_fields_with_valid_values_are_accepted(
    fusion_estate: FusionEstate, records: dict[str, dict[str, Any]], body: dict[str, Any]
) -> None:
    admin = fusion_estate.admin
    activity_id = records["crm"]["activity"]["id"]
    assert (
        admin.patch(f"/api/v1/activities/{activity_id}", json=body, headers=csrf(admin)).status_code
        == 422
    )


# -- review workflow -----------------------------------------------------------------------------


def test_an_incomplete_record_cannot_be_submitted(
    fusion_estate: FusionEstate, records: dict[str, dict[str, Any]]
) -> None:
    analyst = client_for(fusion_estate.app, fusion_estate.users["analyst"])
    activity_id = records["clinic_portal"]["activity"]["id"]
    response = analyst.post(f"/api/v1/activities/{activity_id}/submit", headers=csrf(analyst))
    assert response.status_code == 409
    assert response.json()["detail"]["missing_fields"] == ["special_category_condition"]


def test_submit_review_approve_publish_and_change_after_publication(
    fusion_estate: FusionEstate,
) -> None:
    admin = fusion_estate.admin
    analyst = client_for(fusion_estate.app, fusion_estate.users["analyst"])
    system_id, activity = record_for(fusion_estate, admin, "web_shop", "Shop for review")
    aid = activity["id"]

    assert (
        analyst.post(
            f"/api/v1/activities/{aid}/approve", json={}, headers=csrf(analyst)
        ).status_code
        == 403
    )
    submitted = analyst.post(f"/api/v1/activities/{aid}/submit", headers=csrf(analyst))
    assert submitted.status_code == 200 and submitted.json()["status"] == "awaiting_review"
    assert admin.get(f"/api/v1/systems/{system_id}").json()["status"] == "AWAITING_REVIEW"
    assert (
        analyst.post(f"/api/v1/activities/{aid}/submit", headers=csrf(analyst)).status_code == 409
    )
    assert admin.post(f"/api/v1/activities/{aid}/publish", headers=csrf(admin)).status_code == 409

    queue = analyst.get("/api/v1/review-queue").json()
    assert any(i["kind"] == "approval" and i["activity_id"] == aid for i in queue)

    rejected = admin.post(
        f"/api/v1/activities/{aid}/reject",
        json={"note": "Fix the retention wording"},
        headers=csrf(admin),
    )
    assert rejected.status_code == 200 and rejected.json()["status"] == "draft"
    analyst.post(f"/api/v1/activities/{aid}/submit", headers=csrf(analyst))
    approved = admin.post(
        f"/api/v1/activities/{aid}/approve", json={"note": "Reviewed"}, headers=csrf(admin)
    )
    assert approved.status_code == 200 and approved.json()["status"] == "approved"
    versions = admin.get(f"/api/v1/activities/{aid}/versions").json()
    assert versions[-1]["approved_by"] == str(fusion_estate.users["admin"].id)
    published = admin.post(f"/api/v1/activities/{aid}/publish", headers=csrf(admin))
    assert published.json()["status"] == "published"
    assert admin.get(f"/api/v1/systems/{system_id}").json()["status"] == "PUBLISHED"

    edit = analyst.patch(
        f"/api/v1/activities/{aid}", json={"retention": "5 years for orders"}, headers=csrf(analyst)
    )
    assert (
        edit.json()["status"] == "draft" and edit.json()["current_version"] == 2
    )  # needs review again
    actions = {e["action"] for e in admin.get("/api/v1/audit-log", params={"limit": 500}).json()}
    assert {
        "record.generate",
        "record.submit",
        "record.approve",
        "record.reject",
        "record.publish",
        "record.edit",
    } <= actions


def test_the_submitter_cannot_approve_their_own_record(fusion_estate: FusionEstate) -> None:
    admin = fusion_estate.admin
    system_id, activity = record_for(fusion_estate, admin, "newsletter", "Newsletter self review")
    assert (
        admin.post(f"/api/v1/activities/{activity['id']}/submit", headers=csrf(admin)).status_code
        == 200
    )
    response = admin.post(
        f"/api/v1/activities/{activity['id']}/approve", json={}, headers=csrf(admin)
    )
    assert response.status_code == 409 and "cannot approve" in response.json()["detail"]["message"]


def test_review_queue_lists_incomplete_drafts_and_serious_findings(
    fusion_estate: FusionEstate, records: dict[str, dict[str, Any]]
) -> None:
    analyst = client_for(fusion_estate.app, fusion_estate.users["analyst"])
    queue = analyst.get("/api/v1/review-queue").json()
    kinds = {i["kind"] for i in queue}
    assert {"incomplete", "finding"} <= kinds
    assert any(
        i["kind"] == "incomplete" and "special_category_condition" in i["detail"] for i in queue
    )
    severities = [i["severity"] for i in queue]
    order = {"critical": 0, "high": 1, "medium": 2, "low": 3}
    assert severities == sorted(severities, key=lambda s: order[s])
    assert queue[0]["severity"] == "critical"
    assert (
        client_for(fusion_estate.app, fusion_estate.users["viewer"])
        .get("/api/v1/review-queue")
        .status_code
        == 403
    )


def test_generating_a_record_needs_the_comparison(fusion_estate: FusionEstate) -> None:
    from tests.test_discovery_api import create_system

    admin = fusion_estate.admin
    system_id = create_system(admin, "Not yet", owner=fusion_estate.users["system_owner"].id)
    assert admin.post(f"/api/v1/systems/{system_id}/record", headers=csrf(admin)).status_code == 409


def test_records_are_visible_to_owners_of_their_systems_only(
    fusion_estate: FusionEstate, records: dict[str, dict[str, Any]]
) -> None:
    admin = fusion_estate.admin
    system_id, activity = record_for(
        fusion_estate, admin, "crm", "CRM someone else's", owner=fusion_estate.users["analyst"].id
    )
    owner = client_for(fusion_estate.app, fusion_estate.users["system_owner"])
    ids = {a["id"] for a in owner.get("/api/v1/activities").json()}
    assert activity["id"] not in ids and len(ids) >= 8
    assert owner.get(f"/api/v1/activities/{activity['id']}").status_code == 404
    assert owner.get(f"/api/v1/activities/{activity['id']}/versions").status_code == 404
    assert owner.post(f"/api/v1/systems/{system_id}/record", headers=csrf(owner)).status_code == 403
    assert activity["id"] in {
        a["id"]
        for a in client_for(fusion_estate.app, fusion_estate.users["viewer"])
        .get("/api/v1/activities")
        .json()
    }
    drafts = admin.get("/api/v1/activities", params={"status": "draft"}).json()
    assert drafts and all(a["status"] == "draft" for a in drafts)


# -- exports -------------------------------------------------------------------------------------


def test_the_register_exports_to_a_spreadsheet_with_the_draft_notice(
    fusion_estate: FusionEstate, records: dict[str, dict[str, Any]]
) -> None:
    viewer = client_for(fusion_estate.app, fusion_estate.users["viewer"])
    response = viewer.get("/api/v1/exports/register.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "register-of-processing.xlsx" in response.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Register", "Recipients", "Evidence", "Open findings"]
    register = wb["Register"]
    assert (
        register["A1"].value == "Record of processing activities" and register["A2"].value == NOTICE
    )
    header = [c.value for c in register[5]]
    assert (
        header[:4] == ["Activity", "System", "Status", "Version"]
        and "Article 9 condition" in header
    )
    rows = [[c.value for c in row] for row in register.iter_rows(min_row=6) if row[0].value]
    assert len(rows) >= 8
    names = {r[0] for r in rows}
    assert {"web_shop", "clinic_portal", "hr_onboarding"} <= names
    shop = next(r for r in rows if r[0] == "web_shop")
    assert "Stripe (processor, US)" in shop[header.index("Recipients")]
    assert "7 years" in shop[header.index("Retention")]
    for sheet in ("Recipients", "Evidence", "Open findings"):
        assert wb[sheet]["A1"].value == NOTICE
    assert wb["Evidence"].max_row > 100 and wb["Open findings"].max_row > 10
    assert register.oddFooter.center.text == NOTICE


def test_the_register_exports_to_a_pdf_with_the_notice_on_every_page(
    fusion_estate: FusionEstate, records: dict[str, dict[str, Any]]
) -> None:
    analyst = client_for(fusion_estate.app, fusion_estate.users["analyst"])
    response = analyst.get("/api/v1/exports/register.pdf")
    assert response.status_code == 200 and response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-")
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        assert len(pdf.pages) >= 8
        texts = [page.extract_text() or "" for page in pdf.pages]
    assert all(NOTICE in t for t in texts)
    full = "\n".join(texts)
    assert "Record of processing activities" in full and "web_shop" in full and "Stripe" in full
    assert "Page 1" in texts[0]


def test_only_roles_that_may_read_the_register_can_export_it(fusion_estate: FusionEstate) -> None:
    owner = client_for(fusion_estate.app, fusion_estate.users["system_owner"])
    assert owner.get("/api/v1/exports/register.xlsx").status_code == 403
    assert owner.get("/api/v1/exports/register.pdf").status_code == 403
    anonymous = TestClient(fusion_estate.app, base_url="https://testserver")
    assert anonymous.get("/api/v1/exports/register.pdf").status_code in (401, 403)


def test_an_empty_register_still_exports(
    session_factory: Any, fusion_estate: FusionEstate, other_org: Any
) -> None:
    stranger = client_for(fusion_estate.app, other_org.users["admin"])
    xlsx = stranger.get("/api/v1/exports/register.xlsx")
    assert xlsx.status_code == 200
    assert [c.value for c in load_workbook(io.BytesIO(xlsx.content))["Register"][2]][0] == NOTICE
    pdf = stranger.get("/api/v1/exports/register.pdf")
    with pdfplumber.open(io.BytesIO(pdf.content)) as doc:
        assert "No processing activities" in (doc.pages[0].extract_text() or "")
    assert stranger.get("/api/v1/activities").json() == []


def test_exports_are_audited(fusion_estate: FusionEstate) -> None:
    admin = fusion_estate.admin
    admin.get("/api/v1/exports/register.xlsx")
    entries = admin.get("/api/v1/audit-log", params={"action": "export.register"}).json()
    assert entries and entries[0]["details"]["format"] in {"xlsx", "pdf"}
