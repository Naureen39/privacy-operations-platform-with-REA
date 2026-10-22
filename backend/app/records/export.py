"""The register of processing activities as a spreadsheet and a PDF. Both carry the draft notice."""

import io
from dataclasses import dataclass, field
from datetime import date
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.records.builder import NOTICE, summary_text

TITLE = "Record of processing activities"
COLUMNS = (
    ("Activity", "name"),
    ("System", "system"),
    ("Status", "status"),
    ("Version", "version"),
    ("Purposes", "purposes"),
    ("Data subjects", "data_subjects"),
    ("Data categories", "data_categories"),
    ("Recipients", "recipients"),
    ("Transfers", "transfers"),
    ("Retention", "retention"),
    ("Security measures", "security_measures"),
    ("Lawful basis", "lawful_basis"),
    ("Article 9 condition", "special_category_condition"),
    ("Next review", "next_review_date"),
    ("Sources", "sources"),
)


@dataclass
class RegisterRow:
    name: str
    system: str
    status: str
    version: int
    purposes: list[str]
    data_subjects: list[str]
    data_categories: list[str]
    recipients: list[dict[str, Any]]
    transfers: list[dict[str, Any]]
    retention: str
    security_measures: str
    lawful_basis: str
    special_category_condition: str
    next_review_date: str
    citations: dict[str, list[dict[str, Any]]]
    special_categories: list[str] = field(default_factory=list)
    open_findings: list[dict[str, str]] = field(default_factory=list)


def _cells(row: RegisterRow, labels: dict[str, str]) -> dict[str, str]:
    sources = {(c.get("kind"), c.get("id")) for refs in row.citations.values() for c in refs}
    return {
        "name": row.name,
        "system": row.system,
        "status": row.status.replace("_", " "),
        "version": str(row.version),
        "purposes": "; ".join(row.purposes),
        "data_subjects": ", ".join(row.data_subjects),
        "data_categories": ", ".join(labels.get(c, c) for c in row.data_categories),
        "recipients": "\n".join(r.get("text", r.get("vendor", "")) for r in row.recipients)
        or "None",
        "transfers": "\n".join(t.get("text", "") for t in row.transfers) or "None outside the EEA",
        "retention": row.retention or "",
        "security_measures": row.security_measures or "",
        "lawful_basis": (row.lawful_basis or "").replace("_", " "),
        "special_category_condition": (row.special_category_condition or "").replace("_", " "),
        "next_review_date": row.next_review_date or "",
        "sources": f"{len(sources)} evidence references",
    }


def to_xlsx(rows: list[RegisterRow], labels: dict[str, str], generated: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101
    ws.title = "Register"
    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = NOTICE
    ws["A2"].font = Font(italic=True, color="9C2A00")
    ws["A3"] = f"Generated {generated.isoformat()}"
    header_row = 5
    for index, (title, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=index, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for line, row in enumerate(rows, start=header_row + 1):
        cells = _cells(row, labels)
        for index, (_, key) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=line, column=index, value=cells[key])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, (title, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = 14 if len(title) < 10 else 28
    ws.freeze_panes = f"B{header_row + 1}"
    if ws.oddFooter is not None:
        ws.oddFooter.center.text = NOTICE

    rec = wb.create_sheet("Recipients")
    rec.append([NOTICE])
    rec.append(
        [
            "Activity",
            "Recipient",
            "Role",
            "Country",
            "Purpose",
            "Declared in notice and questionnaire",
        ]
    )
    for row in rows:
        for r in row.recipients:
            rec.append(
                [
                    row.name,
                    r.get("vendor"),
                    r.get("role"),
                    r.get("country"),
                    r.get("purpose"),
                    "yes" if r.get("declared") else "no",
                ]
            )

    cit = wb.create_sheet("Evidence")
    cit.append([NOTICE])
    cit.append(["Activity", "Field", "Source", "Reference", "Detail"])
    for row in rows:
        for fname, refs in sorted(row.citations.items()):
            for ref in refs:
                cit.append([row.name, fname, ref.get("kind"), ref.get("id"), ref.get("detail")])

    fnd = wb.create_sheet("Open findings")
    fnd.append([NOTICE])
    fnd.append(["Activity", "Severity", "Type", "Finding"])
    for row in rows:
        for f in row.open_findings:
            fnd.append([row.name, f["severity"], f["type"], f["title"]])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(rows: list[RegisterRow], labels: dict[str, str], generated: date) -> bytes:
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=8, leading=10)
    notice_style = ParagraphStyle(
        "notice", parent=small, textColor=colors.HexColor("#9C2A00"), fontName="Helvetica-Oblique"
    )

    def p(text: str, style: ParagraphStyle = small) -> Paragraph:
        return Paragraph(escape(text).replace("\n", "<br/>"), style)

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica-Oblique", 7)
        canvas.drawString(15 * mm, 8 * mm, NOTICE)
        canvas.drawRightString(landscape(A4)[0] - 15 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
        title=TITLE,
        author="ProcessLedger",
    )
    story: list[Any] = [
        Paragraph(TITLE, styles["Title"]),
        p(NOTICE, notice_style),
        p(f"Generated {generated.isoformat()}"),
        Spacer(1, 6 * mm),
    ]
    if not rows:
        story.append(p("No processing activities have been recorded yet."))
    for row in rows:
        cells = _cells(row, labels)
        story.append(Paragraph(escape(row.name), styles["Heading2"]))
        story.append(p(summary_text(_as_fields(row), labels, row.special_categories)))
        story.append(Spacer(1, 2 * mm))
        table_rows = [
            [p(title, ParagraphStyle("h", parent=small, fontName="Helvetica-Bold")), p(cells[key])]
            for title, key in COLUMNS
            if key not in ("name",)
        ]
        table = Table(table_rows, colWidths=[40 * mm, 215 * mm], repeatRows=0)
        table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                    ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
                ]
            )
        )
        story += [table, Spacer(1, 6 * mm)]
        if row.open_findings:
            story.append(
                p(
                    "Open findings: "
                    + "; ".join(f"[{f['severity']}] {f['title']}" for f in row.open_findings[:10])
                )
            )
            story.append(Spacer(1, 4 * mm))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def _as_fields(row: RegisterRow) -> dict[str, Any]:
    return {
        "name": row.name,
        "data_categories": row.data_categories,
        "data_subjects": row.data_subjects,
        "recipients": row.recipients,
        "retention": row.retention,
    }
